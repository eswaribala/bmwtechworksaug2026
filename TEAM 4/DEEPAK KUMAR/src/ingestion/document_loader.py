from pathlib import Path
from typing import List
from langchain_core.documents import Document
from src.processing.text_cleaner import clean_text
from src.utils.logging_config import setup_logger

logger = setup_logger("document_loader")

def load_single_document(file_path: Path) -> List[Document]:
    """
    Loads text from a single document file (PDF, TXT, DOCX, CSV)
    and returns a list of LangChain Document objects with standardized page/source metadata.
    """
    file_path = Path(file_path)
    if not file_path.exists():
        logger.error(f"File not found: {file_path}")
        return []

    ext = file_path.suffix.lower()
    file_name = file_path.name
    documents = []

    try:
        if ext == ".pdf":
            try:
                from pypdf import PdfReader
                reader = PdfReader(str(file_path))
                for page_num, page in enumerate(reader.pages, start=1):
                    raw_text = page.extract_text() or ""
                    cleaned = clean_text(raw_text)
                    if cleaned:
                        doc = Document(
                            page_content=cleaned,
                            metadata={
                                "document_id": file_name,
                                "document": file_name,
                                "source": file_name,
                                "page": page_num,
                                "document_type": "PDF"
                            }
                        )
                        documents.append(doc)
            except Exception as e:
                logger.error(f"Error extracting PDF {file_name}: {e}")

        elif ext == ".txt":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
            cleaned = clean_text(content)
            if cleaned:
                documents.append(Document(
                    page_content=cleaned,
                    metadata={
                        "document_id": file_name,
                        "document": file_name,
                        "source": file_name,
                        "page": 1,
                        "document_type": "TXT"
                    }
                ))

        elif ext == ".docx":
            try:
                import docx
                doc_obj = docx.Document(str(file_path))
                full_text = "\n".join([para.text for para in doc_obj.paragraphs if para.text.strip()])
                cleaned = clean_text(full_text)
                if cleaned:
                    documents.append(Document(
                        page_content=cleaned,
                        metadata={
                            "document_id": file_name,
                            "document": file_name,
                            "source": file_name,
                            "page": 1,
                            "document_type": "DOCX"
                        }
                    ))
            except Exception as e:
                logger.error(f"Error parsing DOCX {file_name}: {e}")

        elif ext == ".csv":
            import csv
            lines = []
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                reader = csv.reader(f)
                for row in reader:
                    lines.append(", ".join(row))
            cleaned = clean_text("\n".join(lines))
            if cleaned:
                documents.append(Document(
                    page_content=cleaned,
                    metadata={
                        "document_id": file_name,
                        "document": file_name,
                        "source": file_name,
                        "page": 1,
                        "document_type": "CSV"
                    }
                ))
        else:
            logger.warning(f"Unsupported file extension '{ext}' for file {file_name}")

    except Exception as e:
        logger.error(f"Failed to load document {file_name}: {e}")

    logger.info(f"Loaded {len(documents)} page/section documents from {file_name}")
    return documents


def load_documents_from_dir(dir_path: Path) -> List[Document]:
    """Load all supported documents from a specified directory."""
    dir_path = Path(dir_path)
    if not dir_path.exists() or not dir_path.is_dir():
        logger.warning(f"Directory {dir_path} does not exist.")
        return []

    supported_extensions = {".pdf", ".txt", ".docx", ".csv"}
    all_docs = []

    for file_path in dir_path.glob("*"):
        if file_path.is_file() and file_path.suffix.lower() in supported_extensions:
            docs = load_single_document(file_path)
            all_docs.extend(docs)

    logger.info(f"Total {len(all_docs)} page sections loaded from directory {dir_path}")
    return all_docs
